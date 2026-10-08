"""
test_api.py — the main routes, tested end to end through HTTP against a test database.
"""

import csv
import io
from datetime import datetime

from openpyxl import load_workbook


def test_create_rename_duplicate_delete_form(client):
    form = client.post("/api/forms", json={"title": "Survey"}).json()
    assert form["status"] == "draft"
    assert form["questions"] == []

    renamed = client.patch(f"/api/forms/{form['id']}", json={"title": "Better survey"}).json()
    assert renamed["title"] == "Better survey"

    copy = client.post(f"/api/forms/{form['id']}/duplicate").json()
    assert copy["title"] == "Better survey (copy)"
    assert copy["id"] != form["id"]
    assert copy["public_id"] != form["public_id"]

    assert client.delete(f"/api/forms/{form['id']}").status_code == 204
    assert client.get(f"/api/forms/{form['id']}").status_code == 404
    assert len(client.get("/api/forms").json()) == 1


def test_questions_can_be_inserted_and_reordered(client):
    form = client.post("/api/forms", json={"title": "Order"}).json()
    client.post(f"/api/forms/{form['id']}/questions", json={"type": "short_text"})
    client.post(f"/api/forms/{form['id']}/questions", json={"type": "number"})
    form = client.post(f"/api/forms/{form['id']}/questions", json={"type": "email", "position": 0}).json()
    assert [question["type"] for question in form["questions"]] == ["email", "short_text", "number"]

    ids = [question["id"] for question in form["questions"]]
    new_order = [ids[2], ids[0], ids[1]]
    form = client.put(f"/api/forms/{form['id']}/questions/order", json={"question_ids": new_order}).json()
    assert [question["id"] for question in form["questions"]] == new_order
    assert [question["position"] for question in form["questions"]] == [0, 1, 2]

    # An incomplete list must be refused, not silently applied.
    bad = client.put(f"/api/forms/{form['id']}/questions/order", json={"question_ids": ids[:2]})
    assert bad.status_code == 400


def test_empty_form_cannot_be_published(client):
    form = client.post("/api/forms", json={"title": "Empty"}).json()
    assert client.post(f"/api/forms/{form['id']}/publish").status_code == 400


def test_draft_edits_do_not_reach_respondents_until_republished(client, published_form):
    form_id = published_form["form_id"]
    first_question = published_form["questions"][0]

    edited = client.patch(f"/api/questions/{first_question['id']}", json={"title": "New title"}).json()
    assert edited["has_unpublished_changes"] is True

    still_live = client.get(f"/api/public/forms/{published_form['public_id']}").json()
    assert still_live["questions"][0]["title"] == first_question["title"]

    republished = client.post(f"/api/forms/{form_id}/publish").json()
    assert republished["has_unpublished_changes"] is False
    now_live = client.get(f"/api/public/forms/{published_form['public_id']}").json()
    assert now_live["questions"][0]["title"] == "New title"


def test_unpublished_form_is_not_reachable(client, published_form):
    client.post(f"/api/forms/{published_form['form_id']}/unpublish")
    assert client.get(f"/api/public/forms/{published_form['public_id']}").status_code == 404


def submit(client, public_form, answers):
    token = client.post(f"/api/public/forms/{public_form['public_id']}/responses").json()["token"]
    return token, client.post(f"/api/public/responses/{token}/submit", json={"answers": answers})


def valid_answers(public_form):
    questions = public_form["questions_by_type"]
    return [
        {"question_id": questions["short_text"]["id"], "text": "Ann"},
        {"question_id": questions["email"]["id"], "text": "ann@example.com"},
        {"question_id": questions["multiple_choice"]["id"], "choice_ids": [questions["multiple_choice"]["choices"][1]["id"]]},
        {"question_id": questions["yes_no"]["id"], "boolean": False},
        {"question_id": questions["rating"]["id"], "number": 4},
    ]


def test_invalid_submission_is_rejected_and_nothing_is_stored(client, published_form):
    questions = published_form["questions_by_type"]
    _, response = submit(client, published_form, [{"question_id": questions["email"]["id"], "text": "nope"}])
    assert response.status_code == 422
    failed_question_ids = {error["question_id"] for error in response.json()["errors"]}
    assert failed_question_ids == {questions["short_text"]["id"], questions["email"]["id"]}

    table = client.get(f"/api/forms/{published_form['form_id']}/responses").json()
    assert table["responses"] == []


def test_valid_submission_appears_in_results(client, published_form):
    form_id = published_form["form_id"]
    token, response = submit(client, published_form, valid_answers(published_form))
    assert response.status_code == 201

    # Submitting the same response twice is refused.
    again = client.post(f"/api/public/responses/{token}/submit", json={"answers": valid_answers(published_form)})
    assert again.status_code == 409

    table = client.get(f"/api/forms/{form_id}/responses").json()
    assert len(table["responses"]) == 1
    displays = {answer["question_id"]: answer["display"] for answer in table["responses"][0]["answers"]}
    questions = published_form["questions_by_type"]
    assert displays[questions["short_text"]["id"]] == "Ann"
    assert displays[questions["multiple_choice"]["id"]] == "Green"
    assert displays[questions["yes_no"]["id"]] == "No"
    assert displays[questions["rating"]["id"]] == "4"

    detail = client.get(f"/api/responses/{table['responses'][0]['id']}")
    assert detail.status_code == 200

    assert client.get("/api/forms").json()[0]["response_count"] == 1


def test_summary_counts_choices_and_completion_rate(client, published_form):
    submit(client, published_form, valid_answers(published_form))
    submit(client, published_form, valid_answers(published_form))
    # A third person starts but never submits.
    client.post(f"/api/public/forms/{published_form['public_id']}/responses")

    summary = client.get(f"/api/forms/{published_form['form_id']}/summary").json()
    assert summary["started_count"] == 3
    assert summary["submitted_count"] == 2
    assert summary["completion_rate"] == 66.7

    by_type = {question["type"]: question for question in summary["questions"]}
    assert by_type["multiple_choice"]["buckets"] == [
        {"label": "Red", "count": 0},
        {"label": "Green", "count": 2},
        {"label": "Blue", "count": 0},
    ]
    assert by_type["yes_no"]["buckets"] == [{"label": "Yes", "count": 0}, {"label": "No", "count": 2}]
    assert by_type["rating"]["average"] == 4.0
    assert by_type["short_text"]["recent_texts"] == ["Ann", "Ann"]


def test_deleting_an_answered_question_removes_its_answers_only_on_publish(client, published_form):
    form_id = published_form["form_id"]
    rating = published_form["questions_by_type"]["rating"]
    submit(client, published_form, valid_answers(published_form))

    draft = client.delete(f"/api/questions/{rating['id']}").json()
    assert rating["id"] not in [question["id"] for question in draft["questions"]]
    assert draft["answers_lost_on_publish"] == 1
    # The publish confirmation names the deleted question, not just a number.
    assert draft["removed_on_publish"] == [
        {"kind": "question", "label": rating["title"], "question_title": rating["title"], "answer_count": 1}
    ]

    # Still live for respondents, and the stored answer is still there.
    table = client.get(f"/api/forms/{form_id}/responses").json()
    assert rating["id"] in [answer["question_id"] for answer in table["responses"][0]["answers"]]

    published = client.post(f"/api/forms/{form_id}/publish").json()
    assert published["removed_on_publish"] == []
    table = client.get(f"/api/forms/{form_id}/responses").json()
    assert rating["id"] not in [answer["question_id"] for answer in table["responses"][0]["answers"]]
    assert rating["id"] not in [question["id"] for question in table["questions"]]


def test_type_of_an_answered_question_is_locked(client, published_form):
    short_text = published_form["questions_by_type"]["short_text"]
    submit(client, published_form, valid_answers(published_form))
    response = client.patch(f"/api/questions/{short_text['id']}", json={"type": "long_text"})
    assert response.status_code == 409


def test_deleting_a_picked_choice_waits_for_publish(client, published_form):
    form_id = published_form["form_id"]
    multiple_choice = published_form["questions_by_type"]["multiple_choice"]
    green = multiple_choice["choices"][1]
    submit(client, published_form, valid_answers(published_form))

    draft = client.delete(f"/api/choices/{green['id']}").json()
    draft_question = [question for question in draft["questions"] if question["id"] == multiple_choice["id"]][0]
    assert [choice["label"] for choice in draft_question["choices"]] == ["Red", "Blue"]
    assert draft["answers_lost_on_publish"] == 1
    assert draft["removed_on_publish"] == [
        {"kind": "choice", "label": "Green", "question_title": multiple_choice["title"], "answer_count": 1}
    ]

    # The live form still offers Green, and the stored pick is still shown.
    live = client.get(f"/api/public/forms/{published_form['public_id']}").json()
    live_question = [question for question in live["questions"] if question["id"] == multiple_choice["id"]][0]
    assert [choice["label"] for choice in live_question["choices"]] == ["Red", "Green", "Blue"]

    client.post(f"/api/forms/{form_id}/publish")
    summary = client.get(f"/api/forms/{form_id}/summary").json()
    by_type = {question["type"]: question for question in summary["questions"]}
    assert by_type["multiple_choice"]["buckets"] == [{"label": "Red", "count": 0}, {"label": "Blue", "count": 0}]


def test_csv_export_has_a_header_and_one_row_per_response(client, published_form):
    submit(client, published_form, valid_answers(published_form))
    response = client.get(f"/api/forms/{published_form['form_id']}/responses.csv")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")

    rows = list(csv.reader(io.StringIO(response.text)))
    assert len(rows) == 2
    heading, row = rows
    # Typeform's layout: the response number, the questions, then the response details.
    assert heading[0] == "#"
    assert heading[-4:] == ["Response Type", "Start Date (UTC)", "Submit Date (UTC)", "Ending"]
    assert len(heading) == 1 + len(published_form["questions"]) + 4
    assert "Ann" in row and "ann@example.com" in row and "Green" in row
    assert row[-4] == "completed"
    # Dates are written as "2026-10-08 22:07:35".
    assert datetime.strptime(row[-2], "%Y-%m-%d %H:%M:%S") is not None
    assert row[-1] == published_form["thank_you_title"]


def test_xlsx_export_holds_the_same_table_as_the_csv(client, published_form):
    submit(client, published_form, valid_answers(published_form))
    form_id = published_form["form_id"]

    csv_rows = list(csv.reader(io.StringIO(client.get(f"/api/forms/{form_id}/responses.csv").text)))
    response = client.get(f"/api/forms/{form_id}/responses.xlsx")
    assert response.status_code == 200
    assert response.headers["content-disposition"].endswith('.xlsx"')

    sheet = load_workbook(io.BytesIO(response.content)).active
    xlsx_rows = []
    for row in sheet.iter_rows(values_only=True):
        # Excel has no empty string: an empty cell is read back as None.
        xlsx_rows.append(["" if cell is None else str(cell) for cell in row])
    assert xlsx_rows == csv_rows


def test_export_can_be_limited_to_chosen_responses(client, published_form):
    form_id = published_form["form_id"]
    for _ in range(3):
        submit(client, published_form, valid_answers(published_form))
    table = client.get(f"/api/forms/{form_id}/responses").json()
    all_ids = [response["id"] for response in table["responses"]]
    assert table["ending_title"] == published_form["thank_you_title"]

    response = client.get(f"/api/forms/{form_id}/responses.csv", params={"ids": [all_ids[0], all_ids[2]]})
    rows = list(csv.reader(io.StringIO(response.text)))
    assert [row[0] for row in rows[1:]] == [str(all_ids[0]), str(all_ids[2])]

    # An id that belongs to no response of this form simply matches nothing.
    response = client.get(f"/api/forms/{form_id}/responses.csv", params={"ids": [999999]})
    assert len(list(csv.reader(io.StringIO(response.text)))) == 1


def make_branching_form(client):
    """
    Q1 yes/no -> if "No", jump to Q3 (skipping Q2).  Q2 is required.  Q3 short text.
    Returns (form_id, public form JSON).
    """
    form = client.post("/api/forms", json={"title": "Branching"}).json()
    form_id = form["id"]
    client.post(f"/api/forms/{form_id}/questions", json={"type": "yes_no"})
    client.post(f"/api/forms/{form_id}/questions", json={"type": "short_text"})
    form = client.post(f"/api/forms/{form_id}/questions", json={"type": "short_text"}).json()
    first, second, third = form["questions"]
    client.patch(f"/api/questions/{second['id']}", json={"is_required": True})

    created = client.post(
        f"/api/questions/{first['id']}/logic-jumps",
        json={"operator": "is", "compare_boolean": False, "target_question_id": third["id"]},
    )
    assert created.status_code == 201
    assert len(created.json()["questions"][0]["logic_jumps"]) == 1

    client.post(f"/api/forms/{form_id}/publish")
    return form_id, client.get(f"/api/public/forms/{form['public_id']}").json()


def test_logic_jump_skips_a_required_question(client):
    form_id, public_form = make_branching_form(client)
    first, second, third = public_form["questions"]
    assert first["logic_jumps"][0]["target_question_id"] == third["id"]

    # "No" jumps past the required Q2, so leaving Q2 out is fine...
    _, response = submit(client, public_form, [{"question_id": first["id"], "boolean": False}])
    assert response.status_code == 201

    # ...but "Yes" goes through Q2, which is then demanded.
    _, response = submit(client, public_form, [{"question_id": first["id"], "boolean": True}])
    assert response.status_code == 422
    assert [error["question_id"] for error in response.json()["errors"]] == [second["id"]]


def test_answers_to_skipped_questions_are_not_stored(client):
    form_id, public_form = make_branching_form(client)
    first, second, third = public_form["questions"]
    answers = [
        {"question_id": first["id"], "boolean": False},
        {"question_id": second["id"], "text": "should be ignored"},
        {"question_id": third["id"], "text": "kept"},
    ]
    _, response = submit(client, public_form, answers)
    assert response.status_code == 201

    table = client.get(f"/api/forms/{form_id}/responses").json()
    stored_question_ids = [answer["question_id"] for answer in table["responses"][0]["answers"]]
    assert stored_question_ids == [first["id"], third["id"]]


def test_logic_jump_rules_are_checked(client):
    form = client.post("/api/forms", json={"title": "Rules"}).json()
    client.post(f"/api/forms/{form['id']}/questions", json={"type": "short_text"})
    form = client.post(f"/api/forms/{form['id']}/questions", json={"type": "rating"}).json()
    text_question, rating_question = form["questions"]

    # A text question cannot be compared with a number.
    wrong_operator = client.post(
        f"/api/questions/{text_question['id']}/logic-jumps",
        json={"operator": "greater_than", "compare_number": 3},
    )
    assert wrong_operator.status_code == 400

    # A jump may not go backwards.
    backwards = client.post(
        f"/api/questions/{rating_question['id']}/logic-jumps",
        json={"operator": "always", "target_question_id": text_question["id"]},
    )
    assert backwards.status_code == 400

    # A valid rule can be replaced and deleted.
    created = client.post(
        f"/api/questions/{rating_question['id']}/logic-jumps",
        json={"operator": "less_than", "compare_number": 3},
    ).json()
    rule = created["questions"][1]["logic_jumps"][0]
    assert rule["target_question_id"] is None

    replaced = client.put(f"/api/logic-jumps/{rule['id']}", json={"operator": "is", "compare_number": 5}).json()
    assert replaced["questions"][1]["logic_jumps"][0]["operator"] == "is"

    deleted = client.delete(f"/api/logic-jumps/{rule['id']}").json()
    assert deleted["questions"][1]["logic_jumps"] == []


def test_duplicate_copies_logic_jumps_onto_the_new_questions(client):
    form_id, _ = make_branching_form(client)
    copy = client.post(f"/api/forms/{form_id}/duplicate").json()
    first, _, third = copy["questions"]
    assert first["logic_jumps"][0]["target_question_id"] == third["id"]


def test_duplicate_question_goes_right_after_the_original(client):
    form = client.post("/api/forms", json={"title": "Dup"}).json()
    client.post(f"/api/forms/{form['id']}/questions", json={"type": "multiple_choice"})
    form = client.post(f"/api/forms/{form['id']}/questions", json={"type": "email"}).json()
    multiple_choice = form["questions"][0]
    client.patch(f"/api/choices/{multiple_choice['choices'][0]['id']}", json={"label": "One"})
    client.patch(f"/api/questions/{multiple_choice['id']}", json={"title": "Pick", "is_required": True})

    form = client.post(f"/api/questions/{multiple_choice['id']}/duplicate").json()
    assert [question["type"] for question in form["questions"]] == ["multiple_choice", "multiple_choice", "email"]
    copy = form["questions"][1]
    assert copy["id"] != multiple_choice["id"]
    assert copy["title"] == "Pick" and copy["is_required"] is True
    assert [choice["label"] for choice in copy["choices"]] == ["One"]
    assert copy["choices"][0]["id"] != multiple_choice["choices"][0]["id"]


def test_welcome_screen_is_published_with_the_form(client):
    form = client.post("/api/forms", json={"title": "Welcome"}).json()
    assert form["welcome_enabled"] is False
    client.post(f"/api/forms/{form['id']}/questions", json={"type": "short_text"})

    client.post(f"/api/forms/{form['id']}/publish")
    assert client.get(f"/api/public/forms/{form['public_id']}").json()["welcome"] is None

    updated = client.patch(
        f"/api/forms/{form['id']}",
        json={"welcome_enabled": True, "welcome_title": "Hi", "welcome_text": "Two minutes", "welcome_button_text": "Go"},
    ).json()
    assert updated["welcome_enabled"] is True and updated["has_unpublished_changes"] is True

    # Like every other edit, it reaches respondents only after publishing again.
    assert client.get(f"/api/public/forms/{form['public_id']}").json()["welcome"] is None
    client.post(f"/api/forms/{form['id']}/publish")
    assert client.get(f"/api/public/forms/{form['public_id']}").json()["welcome"] == {
        "title": "Hi",
        "text": "Two minutes",
        "button_text": "Go",
    }


def test_choices_can_be_reordered(client):
    form = client.post("/api/forms", json={"title": "Choices"}).json()
    form = client.post(f"/api/forms/{form['id']}/questions", json={"type": "multiple_choice"}).json()
    question = form["questions"][0]
    client.patch(f"/api/choices/{question['choices'][0]['id']}", json={"label": "One"})
    client.post(f"/api/questions/{question['id']}/choices", json={"label": "Two"})
    form = client.post(f"/api/questions/{question['id']}/choices", json={"label": "Three"}).json()
    ids = [choice["id"] for choice in form["questions"][0]["choices"]]

    form = client.put(f"/api/questions/{question['id']}/choices/order", json={"choice_ids": [ids[2], ids[0], ids[1]]}).json()
    assert [choice["label"] for choice in form["questions"][0]["choices"]] == ["Three", "One", "Two"]

    incomplete = client.put(f"/api/questions/{question['id']}/choices/order", json={"choice_ids": ids[:2]})
    assert incomplete.status_code == 400
